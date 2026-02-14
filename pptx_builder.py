"""PPTX generation from extracted PDF page data.

Reconstructs slides with precise positioning of text, images, and shapes.
"""

import math
from io import BytesIO
from typing import Callable, List, Optional

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Pt

from coord_utils import PT_TO_EMU, bbox_to_position_size, pt_to_emu
from font_mapper import map_font
from models import (
    DrawingPath,
    ImageElement,
    PageData,
    TextBlock,
    TextLine,
)


def build_pptx(
    pages: List[PageData],
    output_path: str,
    progress_callback: Optional[Callable[[int, int], None]] = None,
) -> None:
    """Build a PPTX file from extracted page data.

    Args:
        pages: List of PageData from pdf_extractor.
        output_path: Path to save the PPTX file.
        progress_callback: Called with (current_slide, total_slides).
    """
    if not pages:
        raise ValueError("No pages to convert")

    prs = Presentation()

    # Set slide dimensions from first page
    prs.slide_width = pt_to_emu(pages[0].width)
    prs.slide_height = pt_to_emu(pages[0].height)

    # Use blank layout
    blank_layout = None
    for layout in prs.slide_layouts:
        if layout.name == "Blank":
            blank_layout = layout
            break
    if blank_layout is None:
        blank_layout = prs.slide_layouts[6]

    total = len(pages)
    for idx, page_data in enumerate(pages):
        slide = prs.slides.add_slide(blank_layout)
        _build_slide(slide, page_data)
        if progress_callback:
            progress_callback(idx + 1, total)

    prs.save(output_path)


def _build_slide(slide, page_data: PageData) -> None:
    """Populate a single slide with all extracted content."""
    # Full-page fallback takes priority
    if page_data.fallback_image:
        _add_fullpage_image(slide, page_data)
        return

    # Layer 1: Images (behind everything)
    for img in page_data.images:
        _add_image(slide, img)

    # Layer 2: Vector drawings
    for drawing in page_data.drawings:
        _add_drawing(slide, drawing)

    # Layer 3: Text (on top)
    for block in page_data.text_blocks:
        _add_text_block(slide, block)


def _add_text_block(slide, block: TextBlock) -> None:
    """Add a text block as one textbox per line for precise positioning."""
    for line in block.lines:
        if not line.spans:
            continue

        left, top, width, height = bbox_to_position_size(line.bbox)

        # Add buffer to height to prevent clipping
        height = int(height * 1.2)
        # Ensure minimum width
        width = max(width, pt_to_emu(10))

        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.word_wrap = False
        tf.auto_size = MSO_AUTO_SIZE.NONE

        # Zero margins for exact positioning
        tf.margin_left = 0
        tf.margin_right = 0
        tf.margin_top = 0
        tf.margin_bottom = 0

        para = tf.paragraphs[0]
        para.space_before = Pt(0)
        para.space_after = Pt(0)

        for i, span in enumerate(line.spans):
            if not span.text:
                continue

            if i == 0:
                run = para.runs[0] if para.runs else para.add_run()
            else:
                run = para.add_run()

            run.text = span.text

            font = run.font
            font.name = map_font(span.font_name)
            font.size = Pt(span.font_size)
            font.bold = span.is_bold
            font.italic = span.is_italic

            r, g, b = span.color_rgb
            font.color.rgb = RGBColor(r, g, b)

        # Handle rotated text
        dx, dy = line.direction
        if abs(dx) < 0.99 or abs(dy) > 0.01:
            angle_rad = math.atan2(-dy, dx)
            angle_deg = math.degrees(angle_rad)
            if abs(angle_deg) > 0.5:
                txBox.rotation = angle_deg


def _add_image(slide, img: ImageElement) -> None:
    """Place an image at its exact position."""
    left, top, width, height = bbox_to_position_size(img.bbox)

    try:
        image_stream = BytesIO(img.image_bytes)
        slide.shapes.add_picture(image_stream, left, top, width, height)
    except Exception:
        pass


def _add_drawing(slide, drawing: DrawingPath) -> None:
    """Add a vector drawing to the slide."""
    has_curves = any(
        (item[0] == "c" if isinstance(item, (list, tuple)) and len(item) > 0 else False)
        for item in drawing.items
    )

    if has_curves:
        # Bezier curves are not supported by FreeformBuilder.
        # Render as a rectangle with fill as approximation.
        _add_rect_shape(slide, drawing)
        return

    # Simple rectangle detection
    if len(drawing.items) == 1:
        item = drawing.items[0]
        if isinstance(item, (list, tuple)) and len(item) > 0 and item[0] == "re":
            _add_rect_shape(slide, drawing)
            return

    # Use freeform for line-based paths
    _add_freeform_drawing(slide, drawing)


def _add_rect_shape(slide, drawing: DrawingPath) -> None:
    """Add a rectangle shape."""
    left, top, width, height = bbox_to_position_size(drawing.rect)

    if width < 1 or height < 1:
        return

    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)

    # Apply fill
    if drawing.fill_color:
        shape.fill.solid()
        r, g, b = [max(0, min(255, int(c * 255))) for c in drawing.fill_color]
        shape.fill.fore_color.rgb = RGBColor(r, g, b)
    else:
        shape.fill.background()

    # Apply stroke
    if drawing.stroke_color and drawing.line_width > 0:
        r, g, b = [max(0, min(255, int(c * 255))) for c in drawing.stroke_color]
        shape.line.color.rgb = RGBColor(r, g, b)
        shape.line.width = Pt(drawing.line_width)
    else:
        shape.line.fill.background()


def _add_freeform_drawing(slide, drawing: DrawingPath) -> None:
    """Add a line-based path using FreeformBuilder."""
    x0, y0, x1, y1 = drawing.rect
    w = x1 - x0
    h = y1 - y0
    if w <= 0 or h <= 0:
        return

    try:
        builder = slide.shapes.build_freeform(
            start_x=0, start_y=0, scale=PT_TO_EMU
        )

        first_move = True
        for item in drawing.items:
            if not isinstance(item, (list, tuple)) or len(item) < 2:
                continue

            cmd = item[0]

            if cmd == "l":  # line segment
                p1, p2 = item[1], item[2]
                p1x = getattr(p1, "x", p1[0] if isinstance(p1, (list, tuple)) else 0) - x0
                p1y = getattr(p1, "y", p1[1] if isinstance(p1, (list, tuple)) else 0) - y0
                p2x = getattr(p2, "x", p2[0] if isinstance(p2, (list, tuple)) else 0) - x0
                p2y = getattr(p2, "y", p2[1] if isinstance(p2, (list, tuple)) else 0) - y0

                if first_move:
                    builder.move_to(p1x, p1y)
                    first_move = False
                builder.add_line_segments([(p2x, p2y)], close=False)

            elif cmd == "re":  # rectangle
                rect_obj = item[1]
                if hasattr(rect_obj, "x0"):
                    rx0 = rect_obj.x0 - x0
                    ry0 = rect_obj.y0 - y0
                    rx1 = rect_obj.x1 - x0
                    ry1 = rect_obj.y1 - y0
                else:
                    continue

                if first_move:
                    builder.move_to(rx0, ry0)
                    first_move = False
                builder.add_line_segments(
                    [(rx1, ry0), (rx1, ry1), (rx0, ry1), (rx0, ry0)],
                    close=True,
                )

            elif cmd == "qu":  # quad
                quad = item[1]
                if hasattr(quad, "ul"):
                    pts = [
                        (quad.ul.x - x0, quad.ul.y - y0),
                        (quad.ur.x - x0, quad.ur.y - y0),
                        (quad.lr.x - x0, quad.lr.y - y0),
                        (quad.ll.x - x0, quad.ll.y - y0),
                    ]
                    if first_move:
                        builder.move_to(pts[0][0], pts[0][1])
                        first_move = False
                    builder.add_line_segments(pts[1:], close=True)

        if first_move:
            return

        origin_x = pt_to_emu(x0)
        origin_y = pt_to_emu(y0)
        shape = builder.convert_to_shape(origin_x, origin_y)

        # Apply fill
        if drawing.fill_color:
            shape.fill.solid()
            r, g, b = [max(0, min(255, int(c * 255))) for c in drawing.fill_color]
            shape.fill.fore_color.rgb = RGBColor(r, g, b)
        else:
            shape.fill.background()

        # Apply stroke
        if drawing.stroke_color and drawing.line_width > 0:
            r, g, b = [max(0, min(255, int(c * 255))) for c in drawing.stroke_color]
            shape.line.color.rgb = RGBColor(r, g, b)
            shape.line.width = Pt(drawing.line_width)
        else:
            shape.line.fill.background()

    except Exception:
        # Fall back to rectangle if freeform fails
        _add_rect_shape(slide, drawing)


def _add_fullpage_image(slide, page_data: PageData) -> None:
    """Place a full-page raster image covering the entire slide."""
    image_stream = BytesIO(page_data.fallback_image)
    width = pt_to_emu(page_data.width)
    height = pt_to_emu(page_data.height)
    slide.shapes.add_picture(image_stream, 0, 0, width, height)
